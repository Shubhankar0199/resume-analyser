import json
import os
import io
from typing import Optional
from dotenv import load_dotenv

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
from groq import Groq
from pydantic import BaseModel, Field
from pypdf import PdfReader
from docx import Document


# ============================================================
# CONFIGURATION
# ============================================================

load_dotenv()

GROQ_API_KEY = os.getenv("GROQ_API_KEY")

if not GROQ_API_KEY:
    raise RuntimeError(
        "GROQ_API_KEY environment variable is not configured."
    )

client = Groq(api_key=GROQ_API_KEY)

MODEL = "llama-3.3-70b-versatile"

MAX_RESUME_SIZE = 10 * 1024 * 1024  
MAX_JOB_DESCRIPTION_LENGTH = 30000
MAX_RESUME_TEXT_LENGTH = 50000




app = FastAPI(
    title="Resume Analyzer API",
    description="AI-powered resume and job description analyzer",
    version="1.0.0"
)




app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://127.0.0.1:5173",
        "http://localhost:5174",
        "http://127.0.0.1:5174",
        "http://localhost:5175",
        "http://127.0.0.1:5175",
        "http://localhost:3000",
        "http://127.0.0.1:3000",

        # Production frontend
        "https://resume-analyser-frontend-beta.vercel.app",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)




class JobDescription(BaseModel):
    role: Optional[str] = None

    required_skills: list[str] = Field(
        default_factory=list
    )

    preferred_skills: list[str] = Field(
        default_factory=list
    )

    minimum_experience: Optional[float] = None

    education_requirements: list[str] = Field(
        default_factory=list
    )

    responsibilities: list[str] = Field(
        default_factory=list
    )


class Experience(BaseModel):
    title: Optional[str] = None

    company: Optional[str] = None

    start_date: Optional[str] = None

    end_date: Optional[str] = None

    description: Optional[str] = None


class Resume(BaseModel):
    name: Optional[str] = None

    email: Optional[str] = None

    phone: Optional[str] = None

    total_experience_years: Optional[float] = None

    skills: list[str] = Field(
        default_factory=list
    )

    experiences: list[Experience] = Field(
        default_factory=list
    )

    education: list[str] = Field(
        default_factory=list
    )

    projects: list[str] = Field(
        default_factory=list
    )

    certifications: list[str] = Field(
        default_factory=list
    )


class MatchResult(BaseModel):

    
    score: float = Field(
        ge=0,
        le=100
    )

    
    candidate_name: Optional[str] = None

    
    matching_skills: list[str] = Field(
        default_factory=list
    )

    
    missing_important_skills: list[str] = Field(
        default_factory=list
    )

    
    keyword_suggestions: list[str] = Field(
        default_factory=list
    )

    
    experience_requirement_met: bool

    
    final_verdict: str




def groq_json_call(
    system_prompt: str,
    user_prompt: str
) -> dict:

    try:

        response = client.chat.completions.create(

            model=MODEL,

            messages=[
                {
                    "role": "system",
                    "content": system_prompt
                },
                {
                    "role": "user",
                    "content": user_prompt
                }
            ],

            response_format={
                "type": "json_object"
            },

            temperature=0
        )

    except Exception as exc:

        print(
            f"Groq API error: {repr(exc)}"
        )

        raise RuntimeError(
            "Groq API request failed."
        )


    if not response.choices:

        raise RuntimeError(
            "Groq returned no choices."
        )


    content = response.choices[0].message.content

    if not content:

        raise RuntimeError(
            "Groq returned an empty response."
        )


    try:

        return json.loads(content)

    except json.JSONDecodeError as exc:

        print(
            f"Invalid JSON from Groq: {content}"
        )

        raise RuntimeError(
            "Groq returned invalid JSON."
        ) from exc




def parse_job_description(
    job_description_text: str
) -> JobDescription:

    schema = JobDescription.model_json_schema()


    system_prompt = f"""
You are an expert HR assistant and job description parser.

Analyze the supplied job description and extract structured
information.

Return ONLY JSON matching this schema:

{json.dumps(schema, indent=2)}

Rules:

1. Do not invent information.

2. Extract information only from the supplied
   job description.

3. If the job role cannot be determined, return null.

4. If minimum experience is not mentioned,
   return null.

5. If a list has no information, return [].

6. required_skills must contain skills that are explicitly
   required for the position.

7. preferred_skills must contain skills described as:
   preferred, desirable, optional, nice-to-have,
   bonus, or equivalent.

8. responsibilities must contain actual responsibilities
   mentioned in the job description.

9. Keep skills concise.

10. Preserve important technical terms and technologies.
"""


    user_prompt = f"""
JOB DESCRIPTION:

{job_description_text}
"""


    data = groq_json_call(
        system_prompt,
        user_prompt
    )


    try:

        return JobDescription(
            **data
        )

    except Exception as exc:

        raise RuntimeError(
            f"Invalid job description structure: {exc}"
        )




def parse_resume(
    resume_text: str
) -> Resume:

    schema = Resume.model_json_schema()


    system_prompt = f"""
You are an expert resume parser.

Extract structured information from the supplied resume.

Return ONLY JSON matching this schema:

{json.dumps(schema, indent=2)}

Rules:

1. Do not invent information.

2. If information is unavailable, return null.

3. If a list has no information, return [].

4. Include internships inside experiences.

5. Extract skills from the entire resume.

6. Extract projects separately.

7. Preserve dates when available.

8. Calculate total_experience_years only when it can
   reasonably be determined from the resume.

9. Do not infer skills that are not explicitly supported
   by the resume.

10. Keep technical skills concise.

11. Preserve important technologies, frameworks,
    programming languages, databases, cloud platforms,
    tools, and methodologies.
"""


    user_prompt = f"""
RESUME:

{resume_text}
"""


    data = groq_json_call(
        system_prompt,
        user_prompt
    )


    try:

        return Resume(
            **data
        )

    except Exception as exc:

        raise RuntimeError(
            f"Invalid resume structure: {exc}"
        )




def read_pdf(
    file_bytes: bytes
) -> str:

    reader = PdfReader(
        io.BytesIO(file_bytes)
    )

    text_parts = []


    for page in reader.pages:

        try:

            page_text = page.extract_text()

        except Exception:

            page_text = None


        if page_text:

            text_parts.append(
                page_text
            )


    return "\n".join(
        text_parts
    )




def read_docx(
    file_bytes: bytes
) -> str:

    document = Document(
        io.BytesIO(file_bytes)
    )

    text_parts = []


    

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:

            text_parts.append(
                text
            )


    

    for table in document.tables:

        for row in table.rows:

            for cell in row.cells:

                text = cell.text.strip()

                if text:

                    text_parts.append(
                        text
                    )


    return "\n".join(
        text_parts
    )




def extract_resume_text(
    filename: str,
    file_bytes: bytes
) -> str:

    extension = filename.lower().split(".")[-1]


    if extension == "pdf":

        return read_pdf(
            file_bytes
        )


    if extension == "docx":

        return read_docx(
            file_bytes
        )


    raise ValueError(
        "Unsupported file type. "
        "Only PDF and DOCX files are supported."
    )




def calculate_match(
    job: JobDescription,
    resume: Resume
) -> MatchResult:

    schema = MatchResult.model_json_schema()


    system_prompt = f"""
You are an expert technical recruiter,
ATS resume evaluator, and hiring specialist.

Your task is to compare a candidate's resume against
a specific job description.

Return ONLY JSON matching this schema:

{json.dumps(schema, indent=2)}


============================================================
SCORING
============================================================

Give the candidate a score from 0 to 100.

The score must represent the candidate's overall suitability
for THIS SPECIFIC JOB DESCRIPTION.

Consider:

- Required technical skills
- Preferred skills
- Relevant experience
- Experience duration
- Education requirements
- Responsibilities
- Projects
- Certifications
- Important ATS keywords


============================================================
MATCHING SKILLS
============================================================

matching_skills should contain important skills or keywords
that are present in BOTH:

1. Job description
2. Candidate resume


============================================================
MISSING IMPORTANT SKILLS
============================================================

missing_important_skills should contain important skills
that are explicitly present in the job description but
are not present in the candidate resume.

Prioritize required skills over preferred skills.


============================================================
KEYWORD SUGGESTIONS
============================================================

keyword_suggestions should contain useful ATS keywords
or phrases from the job description that are missing
from the resume.

These suggestions should help improve the candidate's
job-match score.

However:

DO NOT tell the candidate to falsely add a skill.

Only suggest a keyword for inclusion when the candidate's
resume provides reasonable evidence that they may have
experience with that skill.

For example:

Job description:
"Experience with Docker"

Resume:
"Deployed applications using containerized environments"

In this case, Docker may be a reasonable keyword suggestion.

But if the resume contains absolutely no evidence of Docker,
do not tell the candidate to claim Docker experience.


============================================================
EXPERIENCE
============================================================

Determine whether the candidate meets the minimum experience
requirement.

If no minimum experience requirement exists, evaluate
experience based on relevance.


============================================================
EDUCATION
============================================================

Evaluate education requirements when they are explicitly
mentioned in the job description.


============================================================
IMPORTANT RULES
============================================================

1. Do not invent candidate skills.

2. Do not assume one technology automatically means
   the candidate knows another technology.

3. Do not treat related skills as identical unless
   there is reasonable evidence.

4. Required skills are more important than preferred skills.

5. Be honest and conservative when calculating the score.

6. Do not inflate the score.

7. Do not penalize the candidate for skills that are
   genuinely irrelevant to the position.

8. Keep final_verdict concise.

9. keyword_suggestions must contain concise,
   ATS-friendly keywords or phrases.

10. missing_important_skills and keyword_suggestions
    can overlap when appropriate.
"""


    user_prompt = f"""
============================================================
JOB DESCRIPTION
============================================================

{job.model_dump_json(indent=2)}


============================================================
CANDIDATE RESUME
============================================================

{resume.model_dump_json(indent=2)}
"""


    data = groq_json_call(
        system_prompt,
        user_prompt
    )


    try:

        return MatchResult(
            **data
        )

    except Exception as exc:

        raise RuntimeError(
            f"Invalid match result structure: {exc}"
        )




@app.get("/")
def health_check():

    return {
        "status": "ok",
        "service": "Resume Analyzer API",
        "version": "1.0.0"
    }




@app.post("/api/analyze")
async def analyze_resume(
    job_description: str = Form(...),
    resume: UploadFile = File(...)
):

    

    job_description = job_description.strip()


    if not job_description:

        raise HTTPException(
            status_code=400,
            detail="Job description cannot be empty."
        )


    if len(job_description) > MAX_JOB_DESCRIPTION_LENGTH:

        raise HTTPException(
            status_code=400,
            detail=(
                "Job description is too long. "
                "Maximum length is 30,000 characters."
            )
        )


    

    if not resume.filename:

        raise HTTPException(
            status_code=400,
            detail="Resume file is required."
        )


    filename = resume.filename.lower()


    if not (
        filename.endswith(".pdf")
        or filename.endswith(".docx")
    ):

        raise HTTPException(
            status_code=400,
            detail=(
                "Unsupported file type. "
                "Only PDF and DOCX resumes are supported."
            )
        )


    

    try:

        resume_bytes = await resume.read()


        

        if len(resume_bytes) > MAX_RESUME_SIZE:

            raise HTTPException(
                status_code=400,
                detail=(
                    "Resume file must be smaller than 10 MB."
                )
            )


        if len(resume_bytes) == 0:

            raise HTTPException(
                status_code=400,
                detail="Uploaded resume is empty."
            )


        

        resume_text = extract_resume_text(
            resume.filename,
            resume_bytes
        )


    except HTTPException:

        raise


    except Exception as exc:

        print(
            f"Resume extraction error: {repr(exc)}"
        )

        raise HTTPException(
            status_code=400,
            detail=(
                "Could not read the uploaded resume."
            )
        )


  

    resume_text = resume_text.strip()


    if not resume_text:

        raise HTTPException(
            status_code=400,
            detail=(
                "Could not extract text from the resume. "
                "Make sure the PDF/DOCX contains selectable text."
            )
        )


    if len(resume_text) > MAX_RESUME_TEXT_LENGTH:

        resume_text = resume_text[
            :MAX_RESUME_TEXT_LENGTH
        ]



    try:

        

        job = parse_job_description(
            job_description
        )


        

        parsed_resume = parse_resume(
            resume_text
        )


        

        result = calculate_match(
            job,
            parsed_resume
        )


    except Exception as exc:

        print(
            f"AI processing error: {repr(exc)}"
        )

        raise HTTPException(
            status_code=500,
            detail=(
                "AI analysis failed. "
                "Please try again."
            )
        )


    

    return JSONResponse(

        content={

            "success": True,

            "job": job.model_dump(),

            "resume": parsed_resume.model_dump(),

            "match": result.model_dump()

        }
    )
